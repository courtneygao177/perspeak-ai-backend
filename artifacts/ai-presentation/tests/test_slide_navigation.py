"""Regression tests for outline navigation and complete multi-page decks."""
import os
import sys
import tempfile
import unittest
import uuid
from unittest import mock


APP_DIR = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, APP_DIR)

import app as app_module  # noqa: E402


SLIDES = [
    {"page": 1, "title": "Opening", "content": "Welcome", "key_claims": []},
    {"page": 2, "title": "Evidence", "content": "Result", "key_claims": []},
    {"page": 3, "title": "Close", "content": "Summary", "key_claims": []},
]


class TestOutlineNavigation(unittest.TestCase):
    def setUp(self):
        app_module.app.config.update(TESTING=True)
        self.client = app_module.app.test_client()
        self.client.post(
            "/register",
            data={
                "email": f"outline-{uuid.uuid4()}@example.com",
                "password": "password123",
            },
        )
        with self.client.session_transaction() as flask_session:
            flask_session["slide_key"] = app_module._save_slides(SLIDES)
            flask_session["state"] = {
                "current_page": 1,
                "in_qa_mode": False,
                "academic_qa_mode": False,
                "follow_up_round": 0,
                "chat_history": [],
            }
            flask_session["answers"] = []

    def test_outline_jump_updates_backend_and_seals_departing_narration(self):
        response = self.client.post(
            "/x/go-to-slide",
            json={"page": 3, "narration": "This belongs to the opening slide."},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["slide"]["title"], "Close")
        with self.client.session_transaction() as flask_session:
            self.assertEqual(flask_session["state"]["current_page"], 3)
            self.assertEqual(flask_session["answers"][-1]["page"], 1)
            self.assertEqual(flask_session["answers"][-1]["navigation"], "outline_jump")

    def test_outline_jump_is_blocked_during_question_and_answer(self):
        with self.client.session_transaction() as flask_session:
            state = dict(flask_session["state"])
            state["in_qa_mode"] = True
            flask_session["state"] = state
        response = self.client.post("/x/go-to-slide", json={"page": 2})
        self.assertEqual(response.status_code, 409)
        with self.client.session_transaction() as flask_session:
            self.assertEqual(flask_session["state"]["current_page"], 1)

    def test_unknown_page_is_rejected(self):
        response = self.client.post("/x/go-to-slide", json={"page": 99})
        self.assertEqual(response.status_code, 404)

    def test_sidebar_contract_is_clickable(self):
        template_path = os.path.join(APP_DIR, "templates", "sandbox.html")
        with open(template_path, encoding="utf-8") as handle:
            template = handle.read()
        self.assertIn("onclick=\"goToSlide({{ slide.page }})\"", template)
        self.assertIn("/x/go-to-slide", template)
        self.assertIn("cursor: pointer", template)


class TestMultiPageDeckParsing(unittest.TestCase):
    def test_pptx_text_parser_keeps_pages_after_twenty(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "long-deck.pptx")
            presentation = Presentation()
            # The default deck begins with one slide only when explicitly added.
            for page in range(1, 24):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
                box.text_frame.text = f"Slide {page}"
            presentation.save(path)

            parsed = app_module.extract_ppt_images_as_base64(path)

        self.assertEqual(len(parsed), 23)
        self.assertEqual(parsed[-1]["page"], 23)

    def test_office_rendering_prefers_converted_pdf(self):
        import fitz

        with tempfile.TemporaryDirectory() as directory:
            source = os.path.join(directory, "deck.pptx")
            open(source, "wb").close()
            pdf_path = os.path.join(directory, "deck.pdf")
            document = fitz.open()
            page = document.new_page(width=800, height=450)
            page.insert_text((72, 100), "Faithful converted page")
            document.save(pdf_path)
            document.close()

            with mock.patch.object(app_module, "_convert_office_to_pdf", return_value=pdf_path):
                with app_module.app.test_request_context("/x/slide-image/1?thumb=1"):
                    response = app_module._serve_slide_image(source, 1, thumbnail=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "image/jpeg")


if __name__ == "__main__":
    unittest.main()
